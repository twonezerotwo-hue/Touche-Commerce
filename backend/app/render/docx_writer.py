"""DOCX output (python-docx): an editable copy of the same document."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.render.formatting import clean_text, format_cell, format_date, format_money
from app.render.labels import label
from app.render.spec import DocumentSpec

_NUMERIC = {"qty", "money"}


def render_docx(spec: DocumentSpec) -> bytes:
    locale = spec.locale
    document = Document()
    document.core_properties.title = clean_text(f"{spec.title} {spec.number}")
    document.core_properties.author = "TOUCHE COMMERCE"
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    if spec.draft:
        banner = document.add_paragraph()
        run = banner.add_run(
            clean_text(
                f"{label(locale, 'draft')} - "
                f"{spec.notice or label(locale, 'notice.draft_unapproved')}"
            )
        )
        run.bold = True
        run.font.color.rgb = RGBColor(0x8A, 0x5A, 0x00)

    document.add_heading(clean_text(spec.title), level=0)
    document.add_paragraph(
        f"{label(locale, 'number')}: {spec.number}    "
        f"{label(locale, 'date')}: {format_date(spec.issue_date, locale)}"
    )
    if spec.notice and not spec.draft:
        document.add_paragraph(clean_text(spec.notice))

    issuer = document.add_paragraph()
    issuer.add_run(f"{label(locale, 'issuer')}: ").bold = True
    issuer.add_run(clean_text(" | ".join([spec.issuer.name, *spec.issuer.lines])))
    for caption, party in spec.parties:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{caption}: ").bold = True
        paragraph.add_run(clean_text(" | ".join([party.name, *party.lines])))

    if spec.facts:
        table = document.add_table(rows=0, cols=2)
        table.style = "Light List"
        for caption, value in spec.facts:
            cells = table.add_row().cells
            cells[0].text = clean_text(caption)
            cells[1].text = clean_text(value)

    if spec.columns:
        document.add_paragraph()
        table = document.add_table(rows=1, cols=len(spec.columns))
        table.style = "Light Grid Accent 1"
        for index, column in enumerate(spec.columns):
            table.rows[0].cells[index].text = clean_text(column.label)
        for row in spec.rows:
            cells = table.add_row().cells
            for index, column in enumerate(spec.columns):
                cells[index].text = format_cell(column.kind, row.get(column.key), locale)
                if column.kind in _NUMERIC:
                    cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for caption, amount in spec.totals:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.add_run(f"{caption}: {format_money(amount, spec.currency, locale)}").bold = True

    if spec.notes:
        document.add_paragraph().add_run(label(locale, "notes")).bold = True
        document.add_paragraph(clean_text(spec.notes))
    if spec.warnings:
        warning = document.add_paragraph()
        warning.add_run(label(locale, "warnings")).bold = True
        for text in spec.warnings:
            document.add_paragraph(clean_text(text), style="List Bullet")

    document.add_paragraph(clean_text(spec.footer)).runs[0].font.size = Pt(8)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
