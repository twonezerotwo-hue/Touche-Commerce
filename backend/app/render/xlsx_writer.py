"""XLSX output (openpyxl) with real numbers and no formulas.

Text that a spreadsheet could interpret as a formula (leading ``= + - @``) is forced to a
string cell, so a counterparty named ``=HYPERLINK(...)`` stays text.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.render.formatting import clean_text, format_date
from app.render.labels import label
from app.render.spec import DocumentSpec

_FORMULA_LEAD = ("=", "+", "-", "@")
_MONEY_FORMAT = "#,##0.00"
_QTY_FORMAT = "#,##0.####"
_AMBER = "8A5A00"


def _put(sheet: Any, row: int, column: int, value: object) -> Any:
    cell = sheet.cell(row=row, column=column)
    if isinstance(value, str):
        text = clean_text(value)
        cell.value = text
        if text.startswith(_FORMULA_LEAD):
            cell.data_type = "s"  # never a formula
    else:
        cell.value = value
    return cell


def _names(party_name: str, lines: tuple[str, ...]) -> str:
    return " | ".join([party_name, *lines])


def render_xlsx(spec: DocumentSpec) -> bytes:
    locale = spec.locale
    workbook = Workbook()
    sheet: Any = workbook.active
    sheet.title = clean_text(spec.title)[:31]
    row = 1

    if spec.draft:
        banner = spec.notice or label(locale, "notice.draft_unapproved")
        cell = _put(sheet, row, 1, f"{label(locale, 'draft')} - {banner}")
        cell.font = Font(bold=True, color=_AMBER)
        row += 1
    _put(sheet, row, 1, spec.title).font = Font(bold=True, size=14)
    row += 1
    _put(sheet, row, 1, f"{label(locale, 'number')}: {spec.number}")
    _put(sheet, row, 3, f"{label(locale, 'date')}: {format_date(spec.issue_date, locale)}")
    row += 1
    issuer = _names(spec.issuer.name, spec.issuer.lines)
    _put(sheet, row, 1, f"{label(locale, 'issuer')}: {issuer}")
    row += 1
    for caption, party in spec.parties:
        _put(sheet, row, 1, f"{caption}: {_names(party.name, party.lines)}")
        row += 1
    for caption, fact in spec.facts:
        _put(sheet, row, 1, caption).font = Font(bold=True)
        _put(sheet, row, 2, fact)
        row += 1
    row += 1

    header_row = row
    fill = PatternFill("solid", fgColor="14213D")
    for index, column in enumerate(spec.columns, start=1):
        cell = _put(sheet, row, index, column.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        numeric = column.kind in ("qty", "money")
        cell.alignment = Alignment(horizontal="right" if numeric else "left")
    row += 1
    for record in spec.rows:
        for index, column in enumerate(spec.columns, start=1):
            value = record.get(column.key)
            if column.kind == "date" and isinstance(value, (date, datetime)):
                _put(sheet, row, index, format_date(value, locale))
            elif column.kind in ("qty", "money") and isinstance(value, Decimal):
                cell = _put(sheet, row, index, float(value))
                cell.number_format = _MONEY_FORMAT if column.kind == "money" else _QTY_FORMAT
            else:
                _put(sheet, row, index, "" if value is None else str(value))
        row += 1

    for caption, amount in spec.totals:
        _put(sheet, row, max(len(spec.columns) - 1, 1), caption).font = Font(bold=True)
        cell = _put(sheet, row, max(len(spec.columns), 1), float(amount))
        cell.number_format = _MONEY_FORMAT
        cell.font = Font(bold=True)
        row += 1

    if spec.notes:
        row += 1
        _put(sheet, row, 1, f"{label(locale, 'notes')}: {spec.notes}")
        row += 1
    for warning in spec.warnings:
        _put(sheet, row, 1, f"{label(locale, 'warnings')}: {warning}").font = Font(color=_AMBER)
        row += 1

    last = header_row + len(spec.rows)
    for index in range(1, len(spec.columns) + 1):
        longest = max(
            (
                len(str(sheet.cell(row=r, column=index).value or ""))
                for r in range(header_row, last + 1)
            ),
            default=8,
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 48)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
