"""Stage 13A: deterministic, bounded CSV/XLSX import PREVIEW — no database writes.

This parser deliberately does not create counterparties or prices. A real import
commit requires company permissions, per-row idempotency, audit and explicit
human confirmation in a separate stage. No spreadsheet formula is evaluated.
"""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from pydantic import ValidationError

from app.api.v1.schemas.trade import CounterpartyCreateRequest, ProductCreateRequest

ImportKind = Literal["products", "counterparties"]
MAX_BYTES = 2_000_000
MAX_ROWS = 500
MAX_COLS = 24
MAX_UNCOMPRESSED_XLSX = 16_000_000

FIELDS = {
    "products": frozenset({"sku", "name", "unit", "currency", "unit_price"}),
    "counterparties": frozenset(
        {"name", "type", "country", "contact_name", "email", "phone", "notes"}
    ),
}
REQUIRED = {
    "products": frozenset({"sku", "name", "unit", "currency", "unit_price"}),
    "counterparties": frozenset({"name"}),
}
ALIASES = {
    "stok_kodu": "sku",
    "ürün_kodu": "sku",
    "urun_kodu": "sku",
    "ürün_adı": "name",
    "urun_adi": "name",
    "firma_adı": "name",
    "firma_adi": "name",
    "birim": "unit",
    "para_birimi": "currency",
    "döviz": "currency",
    "doviz": "currency",
    "birim_fiyat": "unit_price",
    "fiyat": "unit_price",
    "firma_türü": "type",
    "firma_turu": "type",
    "ülke": "country",
    "ulke": "country",
    "irtibat_kişisi": "contact_name",
    "irtibat_kisisi": "contact_name",
    "e_posta": "email",
    "telefon": "phone",
    "notlar": "notes",
}


class ImportPreviewError(ValueError):
    """File is unsupported, unsafe, oversized or lacks a valid schema."""


@dataclass(frozen=True)
class PreviewRow:
    row: int
    valid: bool
    values: dict[str, str]
    errors: tuple[str, ...]


@dataclass(frozen=True)
class ImportPreview:
    kind: ImportKind
    source_type: str
    total_rows: int
    accepted: int
    rejected: int
    rows: tuple[PreviewRow, ...]
    committed: bool = False


def _normal_header(value: str) -> str:
    return value.strip().casefold().replace(" ", "_").replace("-", "_")


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 1024:
                raise ImportPreviewError("Too many spreadsheet members")
            if sum(x.file_size for x in entries) > MAX_UNCOMPRESSED_XLSX:
                raise ImportPreviewError("Spreadsheet uncompressed size exceeds limit")
            for member in entries:
                if member.file_size > 8_000_000:
                    raise ImportPreviewError("Spreadsheet member exceeds size limit")
                if member.file_size and member.file_size / max(1, member.compress_size) > 120:
                    raise ImportPreviewError("Suspicious spreadsheet compression ratio")
                if not member.filename.startswith(("xl/", "_rels/", "docProps/", "[Content_Types]")):
                    raise ImportPreviewError("Unexpected spreadsheet member")
    except (zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise ImportPreviewError("Invalid XLSX archive") from exc


def _sheet_rows(data: bytes) -> list[list[object]]:
    _check_zip(data)
    try:
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            worksheet = workbook.active
            if worksheet is None:
                raise ImportPreviewError("Spreadsheet contains no active sheet")
            rows: list[list[object]] = []
            for index, cells in enumerate(worksheet.iter_rows(values_only=False)):
                if index > MAX_ROWS:
                    raise ImportPreviewError("More than 500 rows are not permitted")
                if len(cells) > MAX_COLS:
                    raise ImportPreviewError("More than 24 columns are not permitted")
                values: list[object] = []
                for cell in cells:
                    if cell.data_type == "f":
                        raise ImportPreviewError("Spreadsheet formulas cannot be imported")
                    value = cell.value
                    if value is not None and not isinstance(value, (str, int, float, Decimal, bool)):
                        raise ImportPreviewError("Unsupported cell value type")
                    values.append(value)
                rows.append(values)
            return rows
        finally:
            workbook.close()
    except ImportPreviewError:
        raise
    except (ValueError, KeyError, TypeError, OSError, zipfile.BadZipFile) as exc:
        raise ImportPreviewError("Unreadable XLSX worksheet") from exc


def _csv_rows(data: bytes) -> list[list[object]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportPreviewError("CSV must be UTF-8 encoded") from exc
    try:
        source = csv.reader(io.StringIO(text, newline=""), strict=True)
        result = []
        for index, row in enumerate(source):
            if index > MAX_ROWS:
                raise ImportPreviewError("More than 500 rows are not permitted")
            if len(row) > MAX_COLS:
                raise ImportPreviewError("More than 24 columns are not permitted")
            result.append(list(row))
        return result
    except csv.Error as exc:
        raise ImportPreviewError("Invalid CSV quoting or delimiters") from exc


def preview_import(
    filename: str,
    data: bytes,
    *,
    kind: ImportKind,
    mapping: dict[str, str] | None = None,
) -> ImportPreview:
    """Review rows locally without writing data or inferring duplicate DB records."""

    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise ImportPreviewError("Only .csv and .xlsx files are accepted")
    if not data or len(data) > MAX_BYTES:
        raise ImportPreviewError("Import file is empty or exceeds 2MB")
    if kind not in FIELDS:
        raise ImportPreviewError("Unknown import kind")
    if mapping is not None and len(mapping) > MAX_COLS:
        raise ImportPreviewError("Column mapping too large")
    raw_rows = _csv_rows(data) if suffix == ".csv" else _sheet_rows(data)
    if not raw_rows:
        raise ImportPreviewError("No header row")
    headings = [str(value or "") for value in raw_rows[0]]
    columns: dict[str, int] = {}
    for index, name in enumerate(headings):
        normalized = _normal_header(name)
        field = (mapping or {}).get(name, ALIASES.get(normalized, normalized))
        if not field:
            continue
        if field not in FIELDS[kind]:
            raise ImportPreviewError("Unknown column or invalid mapping")
        if field in columns:
            raise ImportPreviewError("Duplicate normalized column")
        columns[field] = index
    if not REQUIRED[kind].issubset(columns):
        raise ImportPreviewError("Required columns are missing")

    model = ProductCreateRequest if kind == "products" else CounterpartyCreateRequest
    observed: set[str] = set()
    results: list[PreviewRow] = []
    for offset, cells in enumerate(raw_rows[1:], start=2):
        record: dict[str, str] = {}
        for field, index in columns.items():
            raw = cells[index] if index < len(cells) else None
            if raw is None:
                continue
            if isinstance(raw, float):
                raw = repr(raw)
            text = str(raw).strip()
            if text.startswith(("=", "+", "-", "@")) and field not in {"unit_price"}:
                raise ImportPreviewError("Spreadsheet/CSV formula-like values are rejected")
            if len(text) > 2100:
                raise ImportPreviewError("Cell text exceeds allowed length")
            if text:
                record[field] = text
        if not record:
            continue
        key = record.get("sku", "") if kind == "products" else record.get("name", "")
        key = key.casefold()
        errors: list[str] = []
        if key in observed:
            errors.append("Duplicate key within file")
        observed.add(key)
        try:
            validated = model.model_validate(record)
            safe = validated.model_dump(mode="json", exclude_none=True)
        except ValidationError as exc:
            errors.extend(
                ".".join(str(part) for part in error["loc"]) + ": invalid value"
                for error in exc.errors()
            )
            safe = record
        results.append(PreviewRow(offset, not errors, safe, tuple(errors)))
    accepted = sum(row.valid for row in results)
    return ImportPreview(kind, suffix.lstrip("."), len(results), accepted,
                         len(results) - accepted, tuple(results))
