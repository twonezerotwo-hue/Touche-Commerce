# ruff: noqa: RUF001
"""Stage 13 document writers: pure functions of a spec, no database."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.render.formatting import (
    clean_text,
    format_date,
    format_decimal,
    format_money,
    format_quantity,
)
from app.render.spec import Column, DocumentSpec, Locale, Party
from app.render.writers import render

COLUMNS = (
    Column("line", "#"),
    Column("sku", "Stok kodu"),
    Column("description", "Açıklama"),
    Column("quantity", "Miktar", "qty"),
    Column("unit_price", "Birim fiyat", "money"),
    Column("line_total", "Tutar", "money"),
)


def spec(
    *,
    locale: Locale = "tr",
    draft: bool = True,
    rows: int = 2,
    description: str = "Şeker pancarı küspesi — İzmir çıkışlı",
) -> DocumentSpec:
    lines = tuple(
        {
            "line": index,
            "sku": f"SKU-{index:03d}",
            "description": description,
            "quantity": Decimal("2500.0000"),
            "unit_price": Decimal("1234.5"),
            "line_total": Decimal("3086250.00"),
        }
        for index in range(1, rows + 1)
    )
    return DocumentSpec(
        kind="sales_quotation",
        locale=locale,
        title="Satış Teklifi" if locale == "tr" else "Sales Quotation",
        number="SQ-2026-0001",
        issue_date=date(2026, 10, 24),
        draft=draft,
        notice="Not yet approved." if locale == "en" else "Henüz onaylanmamıştır.",
        issuer=Party("Touche Acceptance Trading", ("Şişli, İstanbul", "TR")),
        parties=(("Müşteri", Party("Anadolu Gıda Ltd. Şti.", ("ğüşiöç İĞÜŞÖÇ",))),),
        facts=(("Incoterm", "FOB"), ("Ödeme koşulları", "30 gün")),
        columns=COLUMNS,
        rows=lines,
        currency="USD",
        totals=(("Ara toplam", Decimal("6172500.00")), ("Genel toplam", Decimal("6172500.00"))),
        notes="Not: fiyatlar geçerlilik tarihine kadar sabittir.",
        warnings=("Satıcı vergi numarası tanımlı değil.",),
        footer="TOUCHE COMMERCE ile oluşturuldu",
        filename="SQ-2026-0001",
    )


def pdf_text(data: bytes) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)


def test_number_and_date_formatting_follows_the_document_locale() -> None:
    assert format_decimal(Decimal("1234567.5"), "tr") == "1.234.567,50"
    assert format_decimal(Decimal("1234567.5"), "en") == "1,234,567.50"
    assert format_decimal(Decimal("-0.005"), "en") == "-0.01"  # half-up, never banker's rounding
    assert format_money(Decimal("10"), "EUR", "tr") == "10,00 EUR"
    assert format_quantity(Decimal("5.0000"), "en") == "5"
    assert format_quantity(Decimal("2.5000"), "tr") == "2,5"
    assert format_date(date(2026, 10, 4), "tr") == "04.10.2026"
    assert format_date(date(2026, 10, 4), "en") == "04 Oct 2026"
    assert clean_text("Fiyat 5 ₺– net\x00") == "Fiyat 5 TRY- net"


@pytest.mark.parametrize("locale", ["tr", "en"])
def test_pdf_keeps_turkish_letters_and_marks_drafts(locale: Locale) -> None:
    text = pdf_text(render(spec(locale=locale), "pdf"))
    assert "Şeker pancarı küspesi" in text
    assert "ğüşiöç İĞÜŞÖÇ" in text
    assert "SQ-2026-0001" in text
    assert "SKU-002" in text
    assert ("TASLAK" if locale == "tr" else "DRAFT") in text
    assert ("24.10.2026" if locale == "tr" else "24 Oct 2026") in text


def test_a_final_document_carries_no_draft_mark() -> None:
    text = pdf_text(render(spec(draft=False), "pdf"))
    assert "TASLAK" not in text
    assert "DRAFT" not in text


def test_pdf_is_deterministic_and_paginates() -> None:
    one = render(spec(rows=3), "pdf")
    assert one == render(spec(rows=3), "pdf")
    assert one.startswith(b"%PDF-")
    long = render(spec(rows=120), "pdf")
    assert len(PdfReader(io.BytesIO(long)).pages) > 1


def test_pdf_never_prints_the_lira_sign_as_a_missing_glyph() -> None:
    text = pdf_text(render(spec(description="Net 100 ₺"), "pdf"))
    assert "Net 100 TRY" in text


def test_docx_is_an_editable_copy_with_locale_numbers() -> None:
    document = Document(io.BytesIO(render(spec(), "docx")))
    paragraphs = "\n".join(p.text for p in document.paragraphs)
    assert "TASLAK" in paragraphs
    assert "Satış Teklifi" in paragraphs
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "3.086.250,00" in cells
    assert "2.500" in cells
    assert "Şeker pancarı küspesi — İzmir çıkışlı".replace("—", "-") in cells


def test_xlsx_keeps_numbers_numeric_and_text_never_becomes_a_formula() -> None:
    hostile = '=HYPERLINK("http://example.invalid","click")'
    workbook = load_workbook(io.BytesIO(render(spec(description=hostile), "xlsx")))
    sheet = workbook.active
    assert sheet is not None
    assert str(sheet["A1"].value).startswith("TASLAK")
    values = [cell for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert not any(cell.data_type == "f" for cell in values)
    hostile_cells = [c for c in values if isinstance(c.value, str) and c.value.startswith("=")]
    assert hostile_cells and all(c.data_type == "s" for c in hostile_cells)
    numbers = [c.value for c in values if isinstance(c.value, (int, float))]
    assert 3086250.0 in numbers
    assert sheet.freeze_panes is not None


def test_csv_is_excel_friendly_and_defuses_formulas() -> None:
    data = render(spec(description="=cmd|'/c calc'!A1"), "csv")
    assert data.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM so Excel shows Turkish letters
    text = data.decode("utf-8-sig")
    lines = text.split("\r\n")
    assert lines[0] == "#,Stok kodu,Açıklama,Miktar,Birim fiyat,Tutar"
    assert "'=cmd|'/c calc'!A1" in lines[1]
    assert "2500.0000" in lines[1] and "1234.5" in lines[1]  # machine-readable, no grouping
